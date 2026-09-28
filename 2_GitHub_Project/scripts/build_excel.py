"""
Sharayia Stocks - Excel Builder
بناء ملف Excel الشامل (محدث مع دعم ATR ووقف الخسارة)
"""

import os
import pandas as pd
from datetime import datetime
import matplotlib
matplotlib.rcParams['font.family'] = 'DejaVu Sans'
matplotlib.rcParams['axes.unicode_minus'] = False
import matplotlib.pyplot as plt

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

from config import EXCEL_FILE, CHART_FILE, HISTORY_FILE, CAIRO_TZ


# ============================================================
# الرسوم البيانية
# ============================================================

def create_charts(df, output_path):
    fig, axes = plt.subplots(2, 2, figsize=(16, 12))

    # (1) توزيع المصادر
    ax1 = axes[0, 0]
    tier_counts = df["عدد المصادر الخضراء"].value_counts().sort_index(ascending=False)
    colors_tier = ["#D0F0F0", "#D9E1F2", "#E2EFDA", "#FFF2CC"]
    bars = ax1.bar([f"{i} مصادر" for i in tier_counts.index], tier_counts.values,
                   color=colors_tier[:len(tier_counts)], edgecolor="black")
    ax1.set_title("توزيع الأسهم حسب عدد المصادر", fontsize=13, fontweight="bold")
    ax1.set_ylabel("عدد الأسهم")
    for b, v in zip(bars, tier_counts.values):
        ax1.text(b.get_x() + b.get_width()/2, b.get_height() + 0.3, str(v),
                 ha='center', fontsize=11, fontweight="bold")
    ax1.grid(axis="y", alpha=0.3)

    # (2) توزيع الإشارات
    ax2 = axes[0, 1]
    signal_map = {
        "🟢🟢 فرصة ممتازة": "فرصة ممتازة",
        "🟢 قريب من الدعم": "قريب من الدعم",
        "🟡 متوسط": "متوسط",
        "🟠 قريب من المقاومة": "قريب من المقاومة",
        "🔴🔴 قريب جداً من المقاومة": "تحذير مقاومة",
        "⚪ محايد": "محايد",
    }
    signal_counts = df["الإشارة"].value_counts()
    signal_labels = [signal_map.get(s, s) for s in signal_counts.index]
    colors_signal = {
        "فرصة ممتازة": "#00B050", "قريب من الدعم": "#92D050",
        "متوسط": "#FFEB9C", "قريب من المقاومة": "#FFC000",
        "تحذير مقاومة": "#FF0000", "محايد": "#A6A6A6"
    }
    sig_colors = [colors_signal.get(s, "#888888") for s in signal_labels]
    ax2.pie(signal_counts.values, labels=signal_labels, autopct='%1.0f%%',
            colors=sig_colors, startangle=90, wedgeprops={"edgecolor": "black"},
            textprops={'fontsize': 9})
    ax2.set_title("توزيع الإشارات", fontsize=13, fontweight="bold")

    # (3) أعلى 10 قطاعات
    ax3 = axes[1, 0]
    sector_counts = df["التصنيف"].value_counts().head(10)
    ax3.barh(sector_counts.index[::-1], sector_counts.values[::-1],
             color="#4A90E2", edgecolor="black")
    ax3.set_title("أعلى 10 قطاعات", fontsize=13, fontweight="bold")
    ax3.set_xlabel("عدد الأسهم")
    for i, v in enumerate(sector_counts.values[::-1]):
        ax3.text(v + 0.2, i, str(v), va='center', fontsize=10, fontweight="bold")
    ax3.grid(axis="x", alpha=0.3)

    # (4) توزيع الثقة
    ax4 = axes[1, 1]
    trust_map = {
        "🟢🟢 ممتازة": "ممتازة", "🟢 عالية": "عالية",
        "🟡 متوسطة": "متوسطة", "🟠 منخفضة": "منخفضة"
    }
    trust_counts = df["مستوى_الثقة"].value_counts()
    trust_labels = [trust_map.get(s, s) for s in trust_counts.index]
    trust_colors_map = {
        "ممتازة": "#00B050", "عالية": "#92D050",
        "متوسطة": "#FFEB9C", "منخفضة": "#FFC000"
    }
    t_colors = [trust_colors_map.get(s, "#888888") for s in trust_labels]
    ax4.bar(trust_labels, trust_counts.values, color=t_colors, edgecolor="black")
    ax4.set_title("توزيع مستوى الثقة", fontsize=13, fontweight="bold")
    ax4.set_ylabel("عدد الأسهم")
    for i, v in enumerate(trust_counts.values):
        ax4.text(i, v + 0.3, str(v), ha='center', fontsize=11, fontweight="bold")
    ax4.grid(axis="y", alpha=0.3)

    plt.tight_layout()
    plt.savefig(output_path, dpi=100, bbox_inches='tight')
    plt.close()
    return output_path


# ============================================================
# بناء الأوراق
# ============================================================

def build_sector_sheet(df, all_cols):
    rows = []
    sectors = sorted(df["التصنيف"].dropna().unique())
    for sector in sectors:
        sub = df[df["التصنيف"] == sector].copy()
        sub = sub.sort_values("عدد المصادر الخضراء", ascending=False)

        header = {c: "" for c in all_cols}
        header["الرمز"] = f"═══ {sector} ({len(sub)} سهم) ═══"
        rows.append(header)

        for _, r in sub.iterrows():
            rows.append({c: r.get(c, "") for c in all_cols})

        avg_src = sub["عدد المصادر الخضراء"].mean()
        pure = len(sub[sub["عدد المصادر الخضراء"] == 6])
        avg_rr = sub["R/R"].mean() if sub["R/R"].notna().any() else 0
        슐 = len(sub[sub["الإشارة"].str.contains("فرصة ممتازة", na=False)])
        near_sup = len(sub[sub["الإشارة"].str.contains("قريب من الدعم", na=False)])

        summary = {c: "" for c in all_cols}
        summary["الرمز"] = (f"   ↳ متوسط المصادر: {avg_src:.2f} | نقية: {pure} | "
                            f"فرص ممتازة: {슐} | قريب من الدعم: {near_sup} | "
                            f"متوسط R/R: {avg_rr:.2f}")
        rows.append(summary)
        rows.append({c: "" for c in all_cols})
    return pd.DataFrame(rows, columns=all_cols)


def build_tier_sheet(df, all_cols):
    rows = []
    tiers = [
        (6, "✨ نقية — متفقة عليها كل المصادر الستة"),
        (5, "🥈 قوية جدًا — 5 من 6 مصادر"),
        (4, "🥉 قوية — 4 من 6 مصادر"),
        (3, "✅ مقبولة — 3 من 6 مصادر"),
    ]
    for tier_num, tier_name in tiers:
        sub = df[df["عدد المصادر الخضراء"] == tier_num].copy()
        sub = sub.sort_values("R/R", ascending=False)

        header = {c: "" for c in all_cols}
        header["الرمز"] = f"═══ {tier_name} ({len(sub)} سهم) ═══"
        rows.append(header)

        for _, r in sub.iterrows():
            rows.append({c: r.get(c, "") for c in all_cols})

        if len(sub) > 0:
            by_sector = sub["التصنيف"].value_counts().head(3).to_dict()
            sectors_str = " ، ".join([f"{k}({v})" for k, v in by_sector.items()])
            avg_rr = sub["R/R"].mean() if sub["R/R"].notna().any() else 0
            슐 = len(sub[sub["الإشارة"].str.contains("فرصة ممتازة", na=False)])

            summary = {c: "" for c in all_cols}
            summary["الرمز"] = (f"   ↳ أكثر القطاعات: {sectors_str} | "
                                f"فرص ممتازة: {슐} | متوسط R/R: {avg_rr:.2f}")
            rows.append(summary)
        rows.append({c: "" for c in all_cols})
    return pd.DataFrame(rows, columns=all_cols)


# ============================================================
# التنسيق
# ============================================================

def format_excel(excel_path):
    wb = load_workbook(excel_path)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    sector_fill = PatternFill("solid", fgColor="D9E1F2")
    pure_fill = PatternFill("solid", fgColor="D0F0F0")
    silver_fill = PatternFill("solid", fgColor="D9E1F2")
    bronze_fill = PatternFill("solid", fgColor="E2EFDA")
    ok_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    KMB_FORMAT = '[>=1000000000]0.00,,,"B";[>=1000000]0.00,,"M";0.00,"K"'

    for sheet_name in wb.sheetnames:
        if sheet_name in ["Dashboard", "السجل التاريخي", "دليل الاستخدام"]:
            continue

        ws = wb[sheet_name]

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 4, 45)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
            cell.border = border
        ws.row_dimensions[1].height = 30

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="right", vertical="center",
                                           wrap_text=True)

    # تنسيق شرطي على البيانات الكاملة
    if "البيانات الكاملة" in wb.sheetnames:
        ws = wb["البيانات الكاملة"]
        header = [cell.value for cell in ws[1]]
        last_row = ws.max_row

        def col_letter(i):
            return chr(64 + i) if i <= 26 else chr(64 + (i-1)//26) + chr(65 + (i-1)%26)

        for col_name in ["حجم_التداول", "القيمة_السوقية"]:
            if col_name in header:
                idx = header.index(col_name) + 1
                for r in range(2, last_row + 1):
                    ws.cell(row=r, column=idx).number_format = KMB_FORMAT

        for col_name, direction in [("% للدعم 1", "green_low"),
                                     ("% للمقاومة 1", "red_low"),
                                     ("R/R", "green_high")]:
            if col_name in header:
                idx = header.index(col_name) + 1
                L = col_letter(idx)
                if direction == "green_low":
                    ws.conditional_formatting.add(f"{L}2:{L}{last_row}",
                        ColorScaleRule(start_type="num", start_value=0,
                                       start_color="63BE7B",
                                       mid_type="num", mid_value=5, mid_color="FFEB84",
                                       end_type="num", end_value=15, end_color="F8696B"))
                elif direction == "red_low":
                    ws.conditional_formatting.add(f"{L}2:{L}{last_row}",
                        ColorScaleRule(start_type="num", start_value=0,
                                       start_color="F8696B",
                                       mid_type="num", mid_value=5, mid_color="FFEB84",
                                       end_type="num", end_value=15, end_color="63BE7B"))
                elif direction == "green_high":
                    ws.conditional_formatting.add(f"{L}2:{L}{last_row}",
                        ColorScaleRule(start_type="num", start_value=0,
                                       start_color="F8696B",
                                       mid_type="num", mid_value=1, mid_color="FFEB84",
                                       end_type="num", end_value=3, end_color="63BE7B"))

        if "الإشارة" in header:
            idx = header.index("الإشارة") + 1
            L = col_letter(idx)
            rng = f"{L}2:{L}{last_row}"
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'ISNUMBER(SEARCH("فرصة ممتازة",{L}2))'],
                fill=PatternFill("solid", fgColor="00B050"),
                font=Font(color="FFFFFF", bold=True)))
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'ISNUMBER(SEARCH("قريب من الدعم",{L}2))'],
                fill=PatternFill("solid", fgColor="C6EFCE")))
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'ISNUMBER(SEARCH("متوسط",{L}2))'],
                fill=PatternFill("solid", fgColor="FFEB9C")))
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'ISNUMBER(SEARCH("قريب من المقاومة",{L}2))'],
                fill=PatternFill("solid", fgColor="FFC7CE")))
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'ISNUMBER(SEARCH("قريب جداً",{L}2))'],
                fill=PatternFill("solid", fgColor="FF0000"),
                font=Font(color="FFFFFF", bold=True)))

        if "عدد المصادر الخضراء" in header:
            idx = header.index("عدد المصادر الخضراء") + 1
            L = col_letter(idx)
            ws.conditional_formatting.add(f"{L}2:{L}{last_row}",
                ColorScaleRule(start_type="num", start_value=3,
                               start_color="FFEB9C",
                               end_type="num", end_value=6, end_color="00B050"))

        if "حالة_التحديث" in header:
            idx = header.index("حالة_التحديث") + 1
            L = col_letter(idx)
            rng = f"{L}2:{L}{last_row}"
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'ISNUMBER(SEARCH("حديث",{L}2))'],
                fill=PatternFill("solid", fgColor="C6EFCE")))
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'ISNUMBER(SEARCH("متوسط",{L}2))'],
                fill=PatternFill("solid", fgColor="FFEB9C")))
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'ISNUMBER(SEARCH("قديم",{L}2))'],
                fill=PatternFill("solid", fgColor="FFC7CE")))

        if "RSI" in header:
            idx = header.index("RSI") + 1
            L = col_letter(idx)
            ws.conditional_formatting.add(f"{L}2:{L}{last_row}",
                ColorScaleRule(start_type="num", start_value=0,
                               start_color="63BE7B",
                               mid_type="num", mid_value=50, mid_color="FFEB84",
                               end_type="num", end_value=100, end_color="F8696B"))

    for sheet_name, fill_type in [("حسب القطاع", "sector"), ("حسب الفئة", "tier")]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2):
            val = str(row[0].value or "")
            if val.startswith("═══"):
                if fill_type == "sector":
                    fill = sector_fill
                else:
                    if "نقية" in val: fill = pure_fill
                    elif "قوية جدًا" in val: fill = silver_fill
                    elif "قوية" in val: fill = bronze_fill
                    else: fill = ok_fill
                for c in row:
                    c.fill = fill
                    c.font = Font(bold=True, size=12, color="1F4E78")
            elif val.strip().startswith("↳"):
                for c in row:
                    c.font = Font(italic=True, color="595959")

    if "الأسهم النقية" in wb.sheetnames:
        ws = wb["الأسهم النقية"]
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.fill = pure_fill

    wb.save(excel_path)


# ============================================================
# Dashboard
# ============================================================

def build_dashboard(excel_path, df, chart_path):
    wb = load_workbook(excel_path)

    if "Dashboard" in wb.sheetnames:
        del wb["Dashboard"]
    ws = wb.create_sheet("Dashboard", 0)

    ws["A1"] = "📊 لوحة التحليل الفني للأسهم الشرعية"
    ws["A1"].font = Font(bold=True, size=18, color="1F4E78")
    ws.merge_cells("A1:J1")

    stats = [
        ("📅 تاريخ التحليل", datetime.now(CAIRO_TZ).strftime("%Y-%m-%d %H:%M")),
        ("📊 إجمالي الأسهم", len(df)),
        ("✨ نقية (6/6 مصادر)", len(df[df["عدد المصادر الخضراء"] == 6])),
        ("🥈 قوية جدًا (5/6)", len(df[df["عدد المصادر الخضراء"] == 5])),
        ("🟢🟢 فرص ممتازة", len(df[df["الإشارة"] == "🟢🟢 فرصة ممتازة"])),
        ("🟢 قريب من الدعم", len(df[df["الإشارة"].str.contains("قريب من الدعم", na=False)])),
        ("🔴 تحذير مقاومة", len(df[df["الإشارة"].str.contains("مقاومة", na=False)])),
        ("🏭 عدد القطاعات", df["التصنيف"].nunique()),
        ("📈 متوسط R/R", f"{df['R/R'].mean():.2f}" if df['R/R'].notna().any() else "N/A"),
    ]

    for i, (label, value) in enumerate(stats, start=3):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = value
        ws[f"A{i}"].font = Font(bold=True, size=11)
        ws[f"B{i}"].font = Font(size=11, color="1F4E78", bold=True)

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 25

    if os.path.exists(chart_path):
        img = XLImage(chart_path)
        img.width = 900
        img.height = 680
        ws.add_image(img, "A15")

    wb.save(excel_path)


# ============================================================
# ورقة السجل التاريخي والدليل
# ============================================================

def build_history_sheet(excel_path):
    wb = load_workbook(excel_path)
    if "السجل التاريخي" in wb.sheetnames:
        del wb["السجل التاريخي"]
    ws = wb.create_sheet("السجل التاريخي")
    ws["A1"] = "📅 السجل التاريخي"
    ws["A1"].font = Font(bold=True, size=16, color="1F4E78")
    ws.merge_cells("A1:D1")
    ws["A3"] = "📂 المسار:"
    ws["B3"] = HISTORY_FILE
    ws["A4"] = "📅 آخر تحديث:"
    ws["B4"] = datetime.now(CAIRO_TZ).strftime("%Y-%m-%d %H:%M")
    if os.path.exists(HISTORY_FILE):
        try:
            df_hist = pd.read_excel(HISTORY_FILE, header=None)
            for r_idx, row in enumerate(df_hist.itertuples(index=False), start=7):
                for c_idx, val in enumerate(row, start=1):
                    if pd.notna(val):
                        ws.cell(row=r_idx, column=c_idx, value=val)
            ws["A6"] = "📊 محتوى السجل:"
            ws["A6"].font = Font(bold=True, size=12, color="1F4E78")
        except Exception as e:
            ws["A6"] = f"⚠️ لم يتم تحميل السجل: {e}"
    else:
        ws["A6"] = "⚠️ لا يوجد سجل تاريخي بعد"
    wb.save(excel_path)


def build_guide_sheet(excel_path):
    wb = load_workbook(excel_path)
    if "دليل الاستخدام" in wb.sheetnames:
        del wb["دليل الاستخدام"]
    ws = wb.create_sheet("دليل الاستخدام")
    
    # إضافة شرح ATR ووقف الخسارة في الدليل
    guide_data = [
        ["═══ 📖 دليل الاستخدام — Sharayia Stocks ═══", "", ""],
        ["ATR (المدى الحقيقي المتوسط)", "يقيس التذبذب الحقيقي للحركة السعرية", "يُستخدم لحساب المخاطرة"],
        ["وقف_الخسارة", "السعر المقترح للخروج الآمن من الصفقة", "محسوب بناءً على ATR والدعم"],
    ]
    for r_idx, row in enumerate(guide_data, start=1):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    wb.save(excel_path)


# ============================================================
# الوظيفة الرئيسية
# ============================================================

def build_all_excel(df):
    print("📊 Building Excel file...")

    df_full = df.copy().reset_index(drop=True); df_full.index += 1
    df_opp = df.dropna(subset=["% للدعم 1"]).copy().sort_values("% للدعم 1").head(40).reset_index(drop=True); df_opp.index += 1
    df_warn = df.dropna(subset=["% للمقاومة 1"]).copy().sort_values("% للمقاومة 1").head(40).reset_index(drop=True); df_warn.index += 1
    df_rr = df.dropna(subset=["R/R"]).copy(); df_rr = df_rr[df_rr["R/R"] > 0].sort_values("R/R", ascending=False).head(40).reset_index(drop=True); df_rr.index += 1
    df_pure = df[df["عدد المصادر الخضراء"] == 6].copy().reset_index(drop=True); df_pure.index += 1

    ALL_COLS = list(df.columns)
    df_by_sector = build_sector_sheet(df, ALL_COLS)
    df_by_tier = build_tier_sheet(df, ALL_COLS)

    with pd.ExcelWriter(EXCEL_FILE, engine="xlsxwriter") as writer:
        df_full.to_excel(writer, sheet_name="البيانات الكاملة", index=False)
        df_opp.to_excel(writer, sheet_name="فرص الشراء", index=False)
        df_warn.to_excel(writer, sheet_name="تحذير مقاومة", index=False)
        df_rr.to_excel(writer, sheet_name="أفضل R-R", index=False)
        df_pure.to_excel(writer, sheet_name="الأسهم النقية", index=False)
        df_by_sector.to_excel(writer, sheet_name="حسب القطاع", index=False)
        df_by_tier.to_excel(writer, sheet_name="حسب الفئة", index=False)

    create_charts(df, CHART_FILE)
    format_excel(EXCEL_FILE)
    build_dashboard(EXCEL_FILE, df, CHART_FILE)
    build_history_sheet(EXCEL_FILE)
    build_guide_sheet(EXCEL_FILE)

    print(f"✅ Excel saved: {EXCEL_FILE}")
    return EXCEL_FILE
